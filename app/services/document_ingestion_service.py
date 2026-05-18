from __future__ import annotations

import json
import re
from collections import Counter
from datetime import datetime
from pathlib import Path
from typing import Any
from uuid import uuid4

from docx import Document
from pypdf import PdfReader
from pptx import Presentation
from werkzeug.datastructures import FileStorage
from werkzeug.utils import secure_filename


class DocumentIngestionService:
    SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".pptx"}
    STOPWORDS = {
        "yang", "dan", "atau", "dengan", "untuk", "dari", "pada", "dalam", "atas", "bank", "indonesia",
        "pasal", "ayat", "bab", "bagian", "para", "sebagai", "bahwa", "oleh", "terhadap", "serta",
        "adalah", "agar", "akan", "juga", "tidak", "dapat", "lebih", "suatu", "telah", "butir",
    }

    def __init__(self, base_dir: Path, data_dir: Path) -> None:
        self._base_dir = base_dir
        self._data_dir = data_dir
        self._raw_dir = base_dir / "knowledge_base" / "raw"
        self._extracted_dir = base_dir / "knowledge_base" / "extracted"
        self._chunks_dir = base_dir / "knowledge_base" / "chunks"
        self._documents_registry = data_dir / "ingested_documents.json"
        self._knowledge_registry = data_dir / "knowledge_chunks.json"
        self._ensure_storage()

    def list_documents(self) -> list[dict[str, Any]]:
        documents = self._read_json(self._documents_registry, default=[])
        return sorted(documents, key=lambda item: item.get("uploaded_at", ""), reverse=True)

    def sync_raw_documents(self) -> dict[str, Any]:
        documents = self._read_json(self._documents_registry, default=[])
        processed_paths = {
            item.get("raw_path")
            for item in documents
            if item.get("status") == "processed" and item.get("raw_path")
        }
        supported_files = [
            path for path in sorted(self._raw_dir.iterdir())
            if path.is_file() and path.suffix.lower() in self.SUPPORTED_EXTENSIONS
        ]

        synced = 0
        skipped = 0
        failed = 0

        for raw_path in supported_files:
            relative_raw_path = self._to_relative_path(raw_path)
            if relative_raw_path in processed_paths:
                skipped += 1
                continue

            try:
                self._ingest_existing_file(raw_path)
                synced += 1
            except ValueError as exc:
                failed += 1
                self._register_failure(raw_path, str(exc))

        return {
            "synced": synced,
            "skipped": skipped,
            "failed": failed,
            "total": len(supported_files),
        }

    def ingest_upload(self, uploaded_file: FileStorage) -> dict[str, Any]:
        if not uploaded_file or not uploaded_file.filename:
            raise ValueError("File dokumen wajib dipilih.")

        original_name = secure_filename(uploaded_file.filename)
        extension = Path(original_name).suffix.lower()
        if extension not in self.SUPPORTED_EXTENSIONS:
            raise ValueError("Format file harus DOCX, PDF teks, atau PPTX.")

        document_id = uuid4().hex
        timestamp = datetime.utcnow().strftime("%Y%m%d%H%M%S")
        stored_name = f"{timestamp}_{document_id}{extension}"
        raw_path = self._raw_dir / stored_name
        uploaded_file.save(raw_path)

        return self._ingest_existing_file(
            raw_path,
            original_name=original_name,
            document_id=document_id,
            uploaded_at=datetime.utcnow().isoformat(timespec="seconds") + "Z",
            note="Siap dipakai sebagai konteks RAG/Gemini karena teks berhasil diekstrak.",
        )

    def _ingest_existing_file(
        self,
        raw_path: Path,
        *,
        original_name: str | None = None,
        document_id: str | None = None,
        uploaded_at: str | None = None,
        note: str | None = None,
    ) -> dict[str, Any]:
        original_name = secure_filename(original_name or raw_path.name)
        extension = raw_path.suffix.lower()
        if extension not in self.SUPPORTED_EXTENSIONS:
            raise ValueError("Format file harus DOCX, PDF teks, atau PPTX.")

        document_id = document_id or uuid4().hex
        timestamp = datetime.utcnow().strftime("%Y%m%d%H%M%S")

        extracted_text = self._extract_text(raw_path)
        normalized_text = self._normalize_text(extracted_text)
        if len(normalized_text.split()) < 40:
            raise ValueError(
                "Teks tidak cukup terbaca. Untuk PDF, pastikan file bukan hasil scan gambar dan memiliki layer teks yang dapat diekstrak."
            )

        chunks = self._chunk_text(normalized_text)
        extracted_path = self._extracted_dir / f"{timestamp}_{document_id}.txt"
        extracted_path.write_text(normalized_text, encoding="utf-8")

        chunk_payload = self._build_chunk_records(document_id, original_name, extension, chunks)
        chunks_path = self._chunks_dir / f"{timestamp}_{document_id}.json"
        chunks_path.write_text(json.dumps(chunk_payload, ensure_ascii=False, indent=2), encoding="utf-8")

        metadata = {
            "document_id": document_id,
            "original_name": original_name,
            "stored_name": raw_path.name,
            "file_type": extension.lstrip("."),
            "uploaded_at": uploaded_at or datetime.utcnow().isoformat(timespec="seconds") + "Z",
            "raw_path": self._to_relative_path(raw_path),
            "extracted_path": self._to_relative_path(extracted_path),
            "chunks_path": self._to_relative_path(chunks_path),
            "chunk_count": len(chunk_payload),
            "word_count": len(normalized_text.split()),
            "character_count": len(normalized_text),
            "status": "processed",
            "note": note or "Dokumen raw berhasil diproses penuh tanpa pemotongan isi dan siap dipakai sebagai konteks RAG/Gemini.",
        }

        documents = self._read_json(self._documents_registry, default=[])
        documents = [item for item in documents if item.get("raw_path") != metadata["raw_path"]]
        documents.insert(0, metadata)
        self._write_json(self._documents_registry, documents)

        current_chunks = self._read_json(self._knowledge_registry, default=[])
        current_chunks = [item for item in current_chunks if item.get("document_id") != document_id]
        current_chunks.extend(chunk_payload)
        self._write_json(self._knowledge_registry, current_chunks)

        return metadata

    def _register_failure(self, raw_path: Path, reason: str) -> None:
        documents = self._read_json(self._documents_registry, default=[])
        relative_path = self._to_relative_path(raw_path)
        failed_metadata = {
            "document_id": f"failed-{raw_path.stem}",
            "original_name": raw_path.name,
            "stored_name": raw_path.name,
            "file_type": raw_path.suffix.lower().lstrip("."),
            "uploaded_at": datetime.utcnow().isoformat(timespec="seconds") + "Z",
            "raw_path": relative_path,
            "extracted_path": "-",
            "chunks_path": "-",
            "chunk_count": 0,
            "word_count": 0,
            "status": "failed",
            "note": reason,
        }
        documents = [item for item in documents if item.get("raw_path") != relative_path]
        documents.insert(0, failed_metadata)
        self._write_json(self._documents_registry, documents)

    def _ensure_storage(self) -> None:
        self._raw_dir.mkdir(parents=True, exist_ok=True)
        self._extracted_dir.mkdir(parents=True, exist_ok=True)
        self._chunks_dir.mkdir(parents=True, exist_ok=True)
        if not self._documents_registry.exists():
            self._documents_registry.write_text("[]\n", encoding="utf-8")
        if not self._knowledge_registry.exists():
            self._knowledge_registry.write_text("[]\n", encoding="utf-8")

    def _extract_text(self, file_path: Path) -> str:
        extension = file_path.suffix.lower()
        if extension == ".docx":
            return self._extract_docx(file_path)
        if extension == ".pdf":
            return self._extract_pdf(file_path)
        if extension == ".pptx":
            return self._extract_pptx(file_path)
        raise ValueError("Format file belum didukung.")

    def _extract_docx(self, file_path: Path) -> str:
        document = Document(file_path)
        parts: list[str] = []
        for paragraph in document.paragraphs:
            text = paragraph.text.strip()
            if text:
                parts.append(text)
        for table in document.tables:
            for row in table.rows:
                values = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if values:
                    parts.append(" | ".join(values))
        return "\n".join(parts)

    def _extract_pdf(self, file_path: Path) -> str:
        reader = PdfReader(str(file_path))
        parts: list[str] = []
        for index, page in enumerate(reader.pages, start=1):
            text = page.extract_text() or ""
            text = text.strip()
            if text:
                parts.append(f"[Halaman {index}]\n{text}")
        if not parts:
            raise ValueError("PDF tidak memiliki teks yang bisa diekstrak. PDF scan gambar belum didukung.")
        return "\n\n".join(parts)

    def _extract_pptx(self, file_path: Path) -> str:
        presentation = Presentation(str(file_path))
        slides: list[str] = []
        for index, slide in enumerate(presentation.slides, start=1):
            parts: list[str] = []
            for shape in slide.shapes:
                if hasattr(shape, "text"):
                    text = shape.text.strip()
                    if text:
                        parts.append(text)
            if parts:
                slides.append(f"[Slide {index}]\n" + "\n".join(parts))
        return "\n\n".join(slides)

    def _normalize_text(self, text: str) -> str:
        text = text.replace("\xa0", " ")
        text = re.sub(r"\r\n?", "\n", text)
        text = re.sub(r"\n{3,}", "\n\n", text)
        text = re.sub(r"[ \t]{2,}", " ", text)
        return text.strip()

    def _chunk_text(self, text: str, chunk_size: int = 2200) -> list[str]:
        text = text.strip()
        if len(text) <= chunk_size:
            return [text]

        paragraphs = [paragraph.strip() for paragraph in re.split(r"\n{2,}", text) if paragraph.strip()]
        chunks: list[str] = []
        current_parts: list[str] = []

        for paragraph in paragraphs:
            candidate = "\n\n".join([*current_parts, paragraph]).strip() if current_parts else paragraph
            if len(candidate) <= chunk_size:
                current_parts.append(paragraph)
                continue

            if current_parts:
                chunks.append("\n\n".join(current_parts).strip())
                current_parts = []

            if len(paragraph) <= chunk_size:
                current_parts = [paragraph]
                continue

            chunks.extend(self._split_large_block(paragraph, chunk_size))

        if current_parts:
            chunks.append("\n\n".join(current_parts).strip())

        return chunks

    def _split_large_block(self, text: str, chunk_size: int) -> list[str]:
        sentences = re.split(r"(?<=[.!?])\s+", text.strip())
        chunks: list[str] = []
        current = ""

        for sentence in sentences:
            sentence = sentence.strip()
            if not sentence:
                continue

            candidate = f"{current} {sentence}".strip() if current else sentence
            if len(candidate) <= chunk_size:
                current = candidate
                continue

            if current:
                chunks.append(current)
                current = ""

            if len(sentence) <= chunk_size:
                current = sentence
                continue

            start = 0
            sentence_length = len(sentence)
            while start < sentence_length:
                end = min(start + chunk_size, sentence_length)
                if end < sentence_length:
                    pivot = sentence.rfind(" ", start, end)
                    if pivot > start + 20:
                        end = pivot
                piece = sentence[start:end].strip()
                if piece:
                    chunks.append(piece)
                start = end

        if current:
            chunks.append(current)
        return [chunk for chunk in chunks if chunk]

    def _build_chunk_records(
        self,
        document_id: str,
        original_name: str,
        extension: str,
        chunks: list[str],
    ) -> list[dict[str, Any]]:
        return [
            {
                "document_id": document_id,
                "chunk_id": f"{document_id}-{index + 1}",
                "code": original_name,
                "title": f"Dokumen {extension.lstrip('.').upper()} · {original_name}",
                "chapter": "Knowledge Upload",
                "article": f"Chunk {index + 1}",
                "summary": self._summarize_chunk(chunk),
                "clause_text": chunk,
                "keywords": self._extract_keywords(chunk),
                "content": chunk,
                "source_type": extension.lstrip("."),
            }
            for index, chunk in enumerate(chunks)
        ]

    def _summarize_chunk(self, chunk: str, max_length: int = 360) -> str:
        flattened = re.sub(r"\s+", " ", chunk).strip()
        if len(flattened) <= max_length:
            return flattened
        return flattened[: max_length - 3].rstrip() + "..."

    def _extract_keywords(self, text: str, limit: int = 12) -> list[str]:
        tokens = re.findall(r"[A-Za-zÀ-ÿ0-9/.-]{4,}", text.lower())
        filtered = [token for token in tokens if token not in self.STOPWORDS]
        ranked = Counter(filtered)
        return [token for token, _ in ranked.most_common(limit)]

    def _read_json(self, path: Path, default: Any) -> Any:
        if not path.exists():
            return default
        with path.open("r", encoding="utf-8") as handle:
            return json.load(handle)

    def _write_json(self, path: Path, payload: Any) -> None:
        path.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")

    def _to_relative_path(self, path: Path) -> str:
        return str(path.relative_to(self._base_dir)).replace("\\", "/")
